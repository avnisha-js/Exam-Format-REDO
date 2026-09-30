"""Reference-exam conformance notes for a teacher Word file.

This does not format the exam and does not invent academic text.
A conforming exam returns no notes. A non-conforming exam is copied
and given ordinary yellow paragraphs the teacher can read and delete.
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

_HEADING = re.compile(r"^section\s*[-–—]?\s*([A-D])\s*$", re.IGNORECASE)
_CLASS = re.compile(
    r"^(?:class|grade)\s*[:\-–—]?\s*(?:\d{1,2}|xii|xi|x)\s*$",
    re.IGNORECASE,
)
_SUBJECT = re.compile(r"^subject\b", re.IGNORECASE)
_TIME = re.compile(r"^time\b", re.IGNORECASE)
_MARKS = re.compile(r"(?:maximum\s+marks|\bm\s*\.\s*m\s*\.)", re.IGNORECASE)

_HEADER_ROLES = (
    (
        "class",
        _CLASS,
        "Class header is missing.",
        "Add the Class header in the header area, following the Reference Exam.",
    ),
    (
        "subject",
        _SUBJECT,
        "Subject header is missing.",
        "Add the Subject header in the header area, following the Reference Exam.",
    ),
    (
        "time",
        _TIME,
        "Time header is missing.",
        "Add the Time header in the header area, following the Reference Exam.",
    ),
    (
        "marks",
        _MARKS,
        "Maximum marks header is missing.",
        "Add the Maximum marks header in the header area, following the Reference Exam.",
    ),
)


@dataclass(frozen=True)
class ExamError:
    code: str
    error: str
    fix: str
    place: str


def _norm(text: str) -> str:
    return " ".join((text or "").split())


def _section_letters(doc: Document) -> list[str]:
    letters = []
    for paragraph in doc.paragraphs:
        match = _HEADING.match(_norm(paragraph.text))
        if match:
            letters.append(match.group(1).upper())
    return letters


def _before_first_section(doc: Document) -> list[str]:
    texts = []
    for paragraph in doc.paragraphs:
        if _HEADING.match(_norm(paragraph.text)):
            break
        texts.append(paragraph.text)
    return texts


def _has_role(texts: list[str], pattern: re.Pattern[str]) -> bool:
    return any(pattern.search(_norm(text)) for text in texts)


def _place_for_missing_section(present: list[str], letter: str) -> str:
    for later in "ABCD":
        if later > letter and later in present:
            return f"before-section:{later}"
    return "bottom"


def find_nonconformances(teacher: Path, reference: Path) -> list[ExamError]:
    """Header roles and section headings the reference has and the teacher lacks."""
    teacher_doc = Document(str(teacher))
    reference_doc = Document(str(reference))
    teacher_header = _before_first_section(teacher_doc)
    reference_header = _before_first_section(reference_doc)
    errors: list[ExamError] = []
    for code, pattern, error, fix in _HEADER_ROLES:
        if _has_role(reference_header, pattern) and not _has_role(teacher_header, pattern):
            errors.append(ExamError(f"missing_header_{code}", error, fix, "header"))
    teacher_sections = _section_letters(teacher_doc)
    seen_sections: list[str] = []
    for letter in _section_letters(reference_doc):
        if letter in seen_sections:
            continue
        seen_sections.append(letter)
        if letter in teacher_sections:
            continue
        errors.append(
            ExamError(
                f"missing_section_{letter}",
                f"Section {letter} is missing.",
                f"Add Section {letter} in the correct position, following the Reference Exam.",
                _place_for_missing_section(teacher_sections, letter),
            )
        )
    return errors


def _paint(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.bold = True
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _insert_before(anchor, text: str):
    new_p = OxmlElement("w:p")
    anchor._p.addprevious(new_p)
    paragraph = Paragraph(new_p, anchor._parent)
    _paint(paragraph, text)
    return paragraph


def _append(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    _paint(paragraph, text)
    return paragraph


def _lines(errors: list[ExamError]) -> list[str]:
    lines = []
    for item in errors:
        lines.append(f"ERROR: {item.error}")
        lines.append(f"FIX: {item.fix}")
    return lines


def _insert_lines_before(anchor, lines: list[str]) -> None:
    # addprevious always inserts immediately before the anchor, so the
    # first line stays above the lines inserted after it.
    for line in lines:
        _insert_before(anchor, line)


def _first_header_paragraph(doc: Document):
    for paragraph in doc.paragraphs:
        text = _norm(paragraph.text)
        if _CLASS.search(text) or _SUBJECT.search(text) or _TIME.search(text) or _MARKS.search(text):
            return paragraph
    return doc.paragraphs[0] if doc.paragraphs else None


def _section_anchors(doc: Document) -> dict[str, object]:
    anchors = {}
    for paragraph in doc.paragraphs:
        match = _HEADING.match(_norm(paragraph.text))
        if match:
            anchors[match.group(1).upper()] = paragraph
    return anchors


def write_nonconformance_file(teacher: Path, output: Path, errors: list[ExamError]) -> Path:
    """Copy the teacher exam and insert yellow ERROR/FIX paragraphs."""
    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(teacher, output)
    doc = Document(str(output))
    header = [item for item in errors if item.place == "header"]
    bottom = [item for item in errors if item.place == "bottom"]
    if header:
        anchor = _first_header_paragraph(doc)
        if anchor is None:
            bottom = header + bottom
        else:
            _insert_lines_before(anchor, _lines(header))
    anchors = _section_anchors(doc)
    for item in errors:
        if not item.place.startswith("before-section:"):
            continue
        letter = item.place.split(":", 1)[1]
        anchor = anchors.get(letter)
        if anchor is None:
            bottom.append(item)
        else:
            _insert_lines_before(anchor, _lines([item]))
    for item in bottom:
        for line in _lines([item]):
            _append(doc, line)
    doc.save(str(output))
    return output
