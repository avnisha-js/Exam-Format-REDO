"""Render the canonical hierarchy as a printable Word exam.

The hierarchy is consumed as-is. Teacher wording is kept. Only structural
labels, spacing, and the Q6 image placement are presentation.
"""

from __future__ import annotations

import re
import zipfile
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_UNDERLINE
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt

from format.styles import (
    BODY_SIZE,
    EXAM_TITLE_SIZE,
    FONT_NAME,
    INDENT_CHOICE,
    INDENT_STEP,
    INDENT_SYLLABUS,
    MAJOR_SIZE,
    MARGIN,
    MAX_IMAGE_WIDTH,
    PAGE_HEIGHT,
    PAGE_WIDTH,
    SCHOOL_SIZE,
    SECTION_SIZE,
    SPACE_CHOICE_AFTER,
    SPACE_CHOICE_BEFORE,
    SPACE_HEADER_EXAM_AFTER,
    SPACE_HEADER_LAST_AFTER,
    SPACE_HEADER_SCHOOL_AFTER,
    SPACE_MAJOR_AFTER,
    SPACE_MAJOR_BEFORE,
    SPACE_OR_AFTER,
    SPACE_OR_BEFORE,
    SPACE_PASSAGE_AFTER,
    SPACE_PASSAGE_BEFORE,
    SPACE_SECTION_AFTER,
    SPACE_SECTION_BEFORE,
    SPACE_SUB_AFTER,
    SPACE_SUB_BEFORE,
    SPACE_SYLLABUS_HEAD_AFTER,
    SPACE_SYLLABUS_HEAD_BEFORE,
)

_LEADING = re.compile(
    r"^\s*(?:"
    r"\d{1,2}[\.\)]\s*(?:\([A-Za-z]\)[\.\)]?\s*)?"
    r"|(?:xii|xi|x|ix|viii|vii|vi|v|iv|iii|ii|i)[\.\)]\s*"
    r"|\([A-Da-d]\)[\.\)]?\s*"
    r"|[A-Da-d][\.\)]\s*"
    r"|\d{1,2}\s+(?=[A-Za-z])"
    r")"
)
_TRAILING_OR = re.compile(r"\s*\bOR\s*$")
_SECTION = re.compile(r"section\s*[-–—]?\s*([A-D])", re.IGNORECASE)
_RULE = re.compile(r"^[\s_\-–—]+$")
_CHOICE_MARK = re.compile(r"([A-Da-d])\s*[\.\)]")


def strip_leading_marker(text: str) -> str:
    stripped = _LEADING.sub("", text or "", count=1)
    return stripped.strip()


def strip_trailing_or(text: str) -> str:
    return _TRAILING_OR.sub("", text or "").rstrip()


def choice_pieces(text: str) -> list[str]:
    """Split one choice block into option texts, in marker order."""
    found: list[tuple[str, int, int]] = []
    raw = text or ""
    for match in _CHOICE_MARK.finditer(raw):
        letter = match.group(1).lower()
        prev = raw[match.start() - 1] if match.start() else ""
        if not found:
            if letter not in "abcd":
                continue
        elif letter != chr(ord(found[-1][0]) + 1):
            continue
        elif prev.isalpha() and letter == "a":
            continue
        found.append((letter, match.start(), match.end()))
        if letter == "d":
            break
    if not found:
        return [raw.strip()]
    pieces = []
    for index, (_letter, _start, end) in enumerate(found):
        stop = found[index + 1][1] if index + 1 < len(found) else len(raw)
        pieces.append(raw[end:stop].strip())
    return pieces


def _ordered(blocks):
    return sorted(
        blocks,
        key=lambda b: (b.source_order, 0 if b.source_id.startswith("p") else 1, b.source_id),
    )


def _is_rule(text: str) -> bool:
    return bool(_RULE.fullmatch(text or ""))


def _media_name(media_part: str) -> str:
    part = media_part.replace("\\", "/")
    if part.startswith("word/"):
        return part
    return "word/" + part.lstrip("/")


def _image_bytes_and_width(teacher_path: Path, media_part: str):
    name = _media_name(media_part)
    with zipfile.ZipFile(teacher_path) as zf:
        data = zf.read(name)
        root_xml = zf.read("word/document.xml")
        width = _default_width()
    from lxml import etree

    root = etree.fromstring(root_xml)
    ns = {"wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"}
    extent = root.find(".//wp:extent", ns)
    if extent is not None and extent.get("cx"):
        width = Emu(int(extent.get("cx")))
    if width > MAX_IMAGE_WIDTH:
        width = MAX_IMAGE_WIDTH
    return data, width


def _default_width():
    return Inches(3.0)


def _apply_run(paragraph, text: str, size, bold: bool = False, underline: bool = False):
    run = paragraph.add_run(text)
    run.bold = bold
    run.underline = WD_UNDERLINE.SINGLE if underline else WD_UNDERLINE.NONE
    run.font.name = FONT_NAME
    run.font.size = size
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    return run


def _paragraph(doc, text, size, *, before=Pt(0), after=Pt(0), left=0, align=None, bold=False, underline=False):
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = before
    fmt.space_after = after
    fmt.line_spacing = 1.15
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    if left:
        fmt.left_indent = left
    if align is not None:
        paragraph.alignment = align
    if text:
        _apply_run(paragraph, text, size, bold=bold, underline=underline)
    return paragraph


def _prepare_document() -> Document:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = BODY_SIZE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.15
    section = doc.sections[0]
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.left_margin = MARGIN
    section.right_margin = MARGIN
    section.top_margin = MARGIN
    section.bottom_margin = MARGIN
    return doc


def _indent(block, by_id, implicit_major_id):
    extra = 0
    parent = by_id.get(block.parent_id)
    seen = set()
    while parent is not None and parent.source_id not in seen:
        seen.add(parent.source_id)
        if parent.block_type == "branch":
            extra += 1
        parent = by_id.get(parent.parent_id)
    if (
        implicit_major_id
        and block.parent_id == implicit_major_id
        and block.block_type != "branch"
    ):
        extra += 1
    if block.block_type == "choice":
        base = INDENT_CHOICE
    elif block.block_type in {"passage", "subquestion", "continuation", "alternative", "branch"}:
        base = INDENT_STEP
    else:
        base = Inches(0)
    return base + (INDENT_STEP * extra)


def _header_lines(blocks):
    ordered = _ordered(blocks)
    syllabus_ids = {b.source_id for b in ordered if b.block_type == "syllabus"}
    first_syllabus = next((b.source_order for b in ordered if b.block_type == "syllabus"), None)
    lines = []
    for block in ordered:
        if block.block_type != "metadata":
            continue
        if not block.original_text.strip() or _is_rule(block.original_text):
            continue
        if first_syllabus is not None and block.source_order < first_syllabus:
            continue
        if block.source_id in syllabus_ids:
            continue
        lines.append(block.original_text.strip())
    return lines


def _section_label(text: str) -> str:
    match = _SECTION.search(text or "")
    if not match:
        return (text or "").strip()
    return f"Section {match.group(1).upper()}"


def format_exam(blocks, teacher_path: Path, out_path: Path) -> Path:
    by_id = {b.source_id: b for b in blocks}
    doc = _prepare_document()
    header = _header_lines(blocks)
    for index, line in enumerate(header):
        if index == 0:
            size, after = SCHOOL_SIZE, SPACE_HEADER_SCHOOL_AFTER
        elif index == 1:
            size, after = EXAM_TITLE_SIZE, SPACE_HEADER_EXAM_AFTER
        else:
            size, after = BODY_SIZE, Pt(0)
        if index == len(header) - 1:
            after = SPACE_HEADER_LAST_AFTER
        _paragraph(doc, line, size, before=Pt(0), after=after, bold=True)

    syllabus = [b for b in _ordered(blocks) if b.block_type == "syllabus"]
    if syllabus:
        _paragraph(
            doc,
            "SYLLABUS",
            MAJOR_SIZE,
            before=SPACE_SYLLABUS_HEAD_BEFORE,
            after=SPACE_SYLLABUS_HEAD_AFTER,
            bold=True,
        )
        for block in syllabus:
            _paragraph(
                doc,
                block.original_text.strip(),
                BODY_SIZE,
                before=Pt(0),
                after=Pt(1),
                left=INDENT_SYLLABUS,
            )

    implicit_major_id = None
    for block in _ordered(blocks):
        kind = block.block_type
        if kind in {"metadata", "syllabus", "other"}:
            continue
        if kind == "section":
            implicit_major_id = None
            _paragraph(
                doc,
                _section_label(block.original_text),
                SECTION_SIZE,
                before=SPACE_SECTION_BEFORE,
                after=SPACE_SECTION_AFTER,
                bold=True,
                underline=True,
            )
            continue
        if kind == "major_question":
            implicit_major_id = block.source_id if block.branch_carrier else None
            stem = strip_trailing_or(strip_leading_marker(block.original_text))
            label = block.final_label or ""
            q_label = label.split()[0] if label and label != "-" else ""
            paragraph = _paragraph(
                doc,
                "",
                MAJOR_SIZE,
                before=SPACE_MAJOR_BEFORE,
                after=SPACE_MAJOR_AFTER,
            )
            if q_label:
                _apply_run(paragraph, q_label, MAJOR_SIZE, bold=True, underline=True)
                if stem:
                    _apply_run(paragraph, " " + stem, MAJOR_SIZE, bold=False, underline=False)
            elif stem:
                _apply_run(paragraph, stem, MAJOR_SIZE, bold=False, underline=False)
            if block.branch_carrier:
                branch_label = label.split()[1] if len(label.split()) > 1 else "A)"
                _paragraph(
                    doc,
                    branch_label,
                    MAJOR_SIZE,
                    before=Pt(2),
                    after=Pt(2),
                    left=INDENT_STEP,
                    bold=True,
                )
            if _TRAILING_OR.search(block.original_text or ""):
                _paragraph(
                    doc,
                    "OR",
                    MAJOR_SIZE,
                    before=SPACE_OR_BEFORE,
                    after=SPACE_OR_AFTER,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                    bold=True,
                )
            continue
        if kind == "or_marker":
            _paragraph(
                doc,
                "OR",
                MAJOR_SIZE,
                before=SPACE_OR_BEFORE,
                after=SPACE_OR_AFTER,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                bold=True,
            )
            continue
        if kind == "image":
            data, width = _image_bytes_and_width(teacher_path, block.media_part)
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.add_run().add_picture(BytesIO(data), width=width)
            continue
        if kind == "branch":
            passage = strip_leading_marker(block.original_text)
            _paragraph(
                doc,
                block.final_label if block.final_label and block.final_label != "-" else "B)",
                MAJOR_SIZE,
                before=SPACE_MAJOR_BEFORE,
                after=Pt(2),
                left=INDENT_STEP,
                bold=True,
            )
            if passage:
                _paragraph(
                    doc,
                    passage,
                    BODY_SIZE,
                    before=SPACE_PASSAGE_BEFORE,
                    after=SPACE_PASSAGE_AFTER,
                    left=INDENT_STEP + INDENT_STEP,
                )
            continue
        left = _indent(block, by_id, implicit_major_id)
        if kind == "choice":
            labels = (block.final_label or "").split()
            pieces = choice_pieces(block.original_text)
            if len(labels) != len(pieces):
                labels = labels or ["a)"]
                while len(labels) < len(pieces):
                    labels.append(labels[-1])
            for label, piece in zip(labels, pieces):
                _paragraph(
                    doc,
                    f"{label} {piece}".strip(),
                    BODY_SIZE,
                    before=SPACE_CHOICE_BEFORE,
                    after=SPACE_CHOICE_AFTER,
                    left=left,
                )
            continue
        if kind == "subquestion":
            body = strip_leading_marker(block.original_text)
            label = block.final_label if block.final_label and block.final_label != "-" else ""
            shown = f"{label} {body}".strip()
            _paragraph(
                doc,
                shown,
                BODY_SIZE,
                before=SPACE_SUB_BEFORE,
                after=SPACE_SUB_AFTER,
                left=left,
            )
            continue
        if kind in {"passage", "continuation", "alternative"}:
            before = SPACE_PASSAGE_BEFORE if kind == "passage" else SPACE_SUB_BEFORE
            after = SPACE_PASSAGE_AFTER if kind == "passage" else SPACE_SUB_AFTER
            _paragraph(
                doc,
                block.original_text.strip(),
                BODY_SIZE,
                before=before,
                after=after,
                left=left,
            )
            continue
        _paragraph(doc, block.original_text.strip(), BODY_SIZE, left=left)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
