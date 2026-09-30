"""Phase 2 checks for the formatted teacher exam."""

from __future__ import annotations

import hashlib
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from format.format_integrity import check_format
from format.word_formatter import format_exam
from run_core import run_pipeline

TEACHER = ROOT / "input" / "Teacher Exam English 10 Actual.docx"
REFERENCE = ROOT / "input" / "Reference Exam English 10 WORD.docx"
OUTPUT = ROOT / "output" / "Formatted Teacher Exam English 10.docx"


def _norm(text: str) -> str:
    return " ".join(text.split())


def test_formatted_exam():
    result = run_pipeline(TEACHER, REFERENCE, ROOT / "output")
    assert result["ok"], result["problems"]
    format_exam(result["blocks"], TEACHER, OUTPUT)
    assert OUTPUT.exists()
    problems = check_format(result["blocks"], TEACHER, OUTPUT, REFERENCE)
    assert problems == [], problems

    from docx import Document

    doc = Document(str(OUTPUT))
    assert not any(p.text.strip().startswith("ERROR:") for p in doc.paragraphs)
    assert not any(p.text.strip().startswith("FIX:") for p in doc.paragraphs)
    text = "\n".join(p.text for p in doc.paragraphs)
    flat = _norm(text)
    cursor = 0
    for n in range(1, 17):
        token = f"Q{n})"
        pos = flat.find(token, cursor)
        assert pos >= 0, token
        cursor = pos + len(token)
    for letter in "ABCD":
        assert f"Section {letter}" in flat
    assert "1.Read the following" not in flat
    assert "9.(a)" not in flat
    assert "Q9)" in flat and "A)" in flat and "B)" in flat
    assert flat.find("Q9)") < flat.find("A)") < flat.find("He hears the last")
    assert "All day long" in flat
    assert "i) The lesson" in flat or "i) The lesson 'His First Flight'" in flat
    assert "large raindrops" in flat and "Diamond" in flat
    assert "vii)" in flat
    assert "two negative powers" in flat
    assert "Mrs. Pumphrey" in flat or "Mrs.Pumphrey" in flat
    assert "OR" in text
    assert "Write a letter to your friend" in flat
    assert "young seagull to finally fly" in flat
    assert "Book 1: First Flight" in flat
    assert "Jalta Sitara" not in text
    syllabus = next(p.text for p in doc.paragraphs if "First Flight" in p.text)
    assert not syllabus.strip().startswith("Q")

    def is_bold(run):
        return run.bold is True

    def is_underlined(run):
        from docx.enum.text import WD_UNDERLINE
        return run.underline not in (None, False, WD_UNDERLINE.NONE)

    for para in doc.paragraphs:
        runs = [r for r in para.runs if r.text]
        stripped = para.text.strip()
        if not runs:
            continue
        if stripped in {f"Section {letter}" for letter in "ABCD"}:
            assert all(is_bold(r) and is_underlined(r) for r in runs), stripped
        elif stripped == "SYLLABUS" or stripped == "OR" or (len(stripped) == 2 and stripped[0].isalpha() and stripped[1] == ")") or stripped.startswith("Bhagini") or stripped.startswith("Quarterly") or stripped.startswith("Class ") or stripped.startswith("Subject") or stripped.startswith("Time"):
            assert all(is_bold(r) and not is_underlined(r) for r in runs), stripped
        elif stripped.startswith("Q") and stripped.split()[0].endswith(")") and stripped.split()[0][1:-1].isdigit():
            assert runs[0].text == stripped.split()[0]
            assert is_bold(runs[0]) and is_underlined(runs[0])
            assert all(not is_bold(r) and not is_underlined(r) for r in runs[1:])
        else:
            assert all(not is_bold(r) and not is_underlined(r) for r in runs), stripped

    with zipfile.ZipFile(TEACHER) as src, zipfile.ZipFile(OUTPUT) as out:
        teacher_hash = hashlib.sha256(src.read("word/media/image1.jpg")).hexdigest()
        media = [n for n in out.namelist() if n.startswith("word/media/")]
        assert len(media) == 1
        assert hashlib.sha256(out.read(media[0])).hexdigest() == teacher_hash
