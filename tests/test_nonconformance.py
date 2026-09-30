"""A teacher exam that misses a reference header and a reference section."""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from format.nonconformance import _HEADING, _norm, find_nonconformances
from web.app import create_app

TEACHER = ROOT / "input" / "Teacher Exam English 10 Actual.docx"
BAD = ROOT / "input" / "Teacher Exam English 10 BAD.docx"
REFERENCE = ROOT / "input" / "Reference Exam English 10 WORD.docx"
OUTPUT = ROOT / "output" / "Teacher Exam English 10 BAD ERRORS.docx"


def _section_letters(doc: Document) -> list[str]:
    letters = []
    for paragraph in doc.paragraphs:
        match = _HEADING.match(_norm(paragraph.text))
        if match:
            letters.append(match.group(1).upper())
    return letters


def test_good_exam_has_no_reference_gaps():
    assert find_nonconformances(TEACHER, REFERENCE) == []


def test_bad_exam_returns_two_yellow_notes(tmp_path):
    app = create_app(tmp_path / "data")
    client = app.test_client()
    uploaded = client.post(
        "/reference",
        data={"reference": (REFERENCE.open("rb"), REFERENCE.name)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Reference Exam English 10 WORD.docx" in uploaded.data

    page = client.post(
        "/format",
        data={"teacher": (BAD.open("rb"), BAD.name)},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    body = page.data.decode("utf-8")
    assert "Formatting failed" not in body
    assert "Teacher Exam English 10 BAD_CORRECTIONS.docx" in body
    start = body.find("/exams/")
    exam_id = body[start + len("/exams/") : body.find("/download", start)]
    download = client.get(f"/exams/{exam_id}/download")
    assert download.status_code == 200
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_bytes(download.data)
    download.close()

    doc = Document(str(OUTPUT))
    errors = [p for p in doc.paragraphs if p.text.strip().startswith("ERROR:")]
    fixes = [p for p in doc.paragraphs if p.text.strip().startswith("FIX:")]
    assert len(errors) == 2
    assert len(fixes) == 2
    error_text = [p.text.strip() for p in errors]
    fix_text = [p.text.strip() for p in fixes]
    assert any("Class header is missing." in text for text in error_text)
    assert any("Section C is missing." in text for text in error_text)
    assert any(text.startswith("FIX: Add the Class header") for text in fix_text)
    assert any(text.startswith("FIX: Add Section C in the correct position") for text in fix_text)
    assert "C" not in _section_letters(doc)
    assert "Jalta Sitara" not in "\n".join(p.text for p in doc.paragraphs)
    assert any("Fill in the blanks" in p.text for p in doc.paragraphs)

    for paragraph in errors + fixes:
        runs = [run for run in paragraph.runs if run.text]
        assert runs
        assert all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in runs)
        assert paragraph._p.xpath(".//*[local-name()='commentReference']") == []

    with zipfile.ZipFile(OUTPUT) as zf:
        names = zf.namelist()
    assert "word/comments.xml" not in names

    broken = client.post(
        "/format",
        data={"teacher": (io.BytesIO(b"this is not a word file"), "broken.docx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"Formatting failed. The exam was not saved." in broken.data
